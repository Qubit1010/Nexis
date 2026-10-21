"""Turn the scored Maps rows into the Islamabad prospect workbook for the three launch sectors.

Reads data/prospects.csv (or prospects_latest.csv when the main file was locked), writes
deliverables/Islamabad-Prospects.xlsx: a summary sheet, one ranked sheet per sector with a
plain-language opportunity line, a suggested Growth bundle, and blank outreach-tracking columns.

Usage:  python scripts/build_prospects.py
"""
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

PROJECT = Path(__file__).resolve().parents[1]
DATA = PROJECT / "data"
OUT = PROJECT / "deliverables"

# analyze_maps sector -> launch sector; everything else is out of scope for this kit
LAUNCH = {"healthcare": "healthcare", "aesthetics": "healthcare", "food": "food", "real-estate": "real-estate"}
AREA_NAME = {"core": "Islamabad core", "dha": "DHA", "bahria": "Bahria Town", "gulberg": "Gulberg", "b17": "B-17"}
# Google categories that a launch-sector search returned but that are not buyers for this kit
OFF_CATEGORY = {"Apartment building", "Shopping mall", "Marketing agency", "Marketing consultant", "E commerce agency",
                "Building materials store", "Architect", "Hotel", "Family law attorney", "Business center",
                "Housing authority", "Association / Organization", "Construction company", "Pharmaceutical company",
                "Animal hospital", "Medical school", "Pharmacy", "Beauty school", "Soup kitchen", "Conference center", ""}
STATUS = ["Not contacted", "Visited", "Messaged", "Health Check booked", "Demo shown", "Proposal sent",
          "Pilot signed", "Not a fit"]

BLACK, BLUE, PAPER = "000000", "02A1E1", "F4F4F2"


def load() -> list[dict]:
    src = DATA / "prospects_latest.csv"
    if not src.exists() or (DATA / "prospects.csv").stat().st_mtime > src.stat().st_mtime:
        src = DATA / "prospects.csv"
    with open(src, encoding="utf-8-sig") as fh:
        rows = [r for r in csv.DictReader(fh) if r["sector"] in LAUNCH
                and not (r["gmaps_category"] in OFF_CATEGORY
                         and not (r["sector"] == "real-estate" and r["gmaps_category"] == "Construction company"))]
    for r in rows:
        r["launch"] = LAUNCH[r["sector"]]
        r["score"] = int(r["score"])
        r["reviews"] = int(float(r["reviews"] or 0))
        r["rating"] = float(r["rating"]) if r["rating"] else None
    return rows


def opportunity(r: dict) -> str:
    notes, status, rev, rat = [], r["website_status"], r["reviews"], r["rating"]
    if status == "none":
        notes.append(f"{rev:,} Google reviews and no website: demand is visible, capture is missing"
                     if rev >= 50 else "No website")
    elif status == "social-only":
        notes.append("Only an Instagram or Facebook page, no site of its own")
    elif status.startswith(("broken", "down")):
        notes.append(f"Website listed but {status}")
    elif status == "live (no https)":
        notes.append("Website has no HTTPS: browsers mark it 'Not secure'")
    if r["listing_claimed"] == "no":
        notes.append("Google listing unclaimed")
    if rat is not None and rev >= 20 and rat < 4.0:
        notes.append(f"Rated {rat} on {rev:,} reviews")
    if not notes and rev >= 100:
        notes.append(f"Busy and online ({rev:,} reviews): sell reply speed and rebooking, not presence")
    return "; ".join(notes) or "Listing in good shape: lead with reply speed"


def opener(r: dict) -> str:
    """Which door-opener fits: the Lead Response Check for property, the Health Check otherwise."""
    if r["launch"] == "real-estate":
        return "Lead Response Check (test enquiry, time the reply)"
    if r["website_status"] in ("none", "social-only") or r["listing_claimed"] == "no":
        return "Digital Health Check (presence gaps)"
    return "Digital Health Check (reply speed + reviews)"


def main() -> None:
    kit = json.loads((DATA / "kit.json").read_text(encoding="utf-8"))
    sectors = {s["key"]: s for s in kit["sectors"]}
    rows = load()
    OUT.mkdir(exist_ok=True)
    wb = Workbook()

    head_font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    head_fill = PatternFill("solid", fgColor=BLACK)
    top_fill = PatternFill("solid", fgColor="E6F6FD")
    wrap = Alignment(wrap_text=True, vertical="top")

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"NexusPoint · Islamabad prospects · Google Maps scan {kit['meta']['date']}"
    ws["A1"].font = Font(name="Calibri", bold=True, size=14)
    ws["A2"] = ("Score = visible demand (reviews) + digital gap + reachability (phone, mobile), out of 100. "
                "Chain branches are halved: decisions sit at head office. Ranked within each sector.")
    ws["A2"].alignment = wrap
    ws.merge_cells("A2:H2")
    ws.row_dimensions[2].height = 32
    hdr = ["Sector", "Businesses", "Score 70+", "No working site", "Listing unclaimed", "Mobile number",
           "Median reviews", "Growth tier (setup + monthly)"]
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font, cell.fill = head_font, head_fill
    for i, key in enumerate(("healthcare", "food", "real-estate"), 5):
        grp = [r for r in rows if r["launch"] == key]
        n = len(grp) or 1
        reviews = sorted(r["reviews"] for r in grp)
        growth = next(t for t in sectors[key]["tiers"] if t.get("rec"))
        vals = [sectors[key]["name"], len(grp), sum(r["score"] >= 70 for r in grp),
                f"{round(100 * sum(r['website_status'] in ('none', 'social-only') or r['website_status'].startswith(('broken', 'down')) for r in grp) / n)}%",
                f"{round(100 * sum(r['listing_claimed'] == 'no' for r in grp) / n)}%",
                f"{round(100 * sum((r['phone'] or '').replace(' ', '').startswith(('+923', '03')) for r in grp) / n)}%",
                reviews[len(reviews) // 2] if reviews else 0,
                f"PKR {growth['setup']:,} + {growth['monthly']:,}/mo"]
        for c, v in enumerate(vals, 1):
            ws.cell(row=i, column=c, value=v)
    ws["A9"] = "Areas: Islamabad core (F, G, E, H, I sectors, Blue Area), DHA, Bahria Town, Gulberg, B-17."
    ws["A10"] = "Status column has a dropdown. Sort or filter freely; the rank column keeps the original order."
    for c, wdt in enumerate([26, 12, 11, 16, 17, 15, 15, 30], 1):
        ws.column_dimensions[get_column_letter(c)].width = wdt

    cols = [("Rank", 6), ("Score", 7), ("Business", 34), ("Type", 22), ("Area", 14), ("Neighbourhood", 20),
            ("Rating", 7), ("Reviews", 9), ("Phone", 17), ("Website", 30), ("Site status", 15), ("Claimed", 9),
            ("Opportunity", 46), ("Door-opener", 30), ("Pitch (Growth)", 30), ("Maps", 12),
            ("Status", 18), ("Owner / contact", 22), ("Last touch", 12), ("Notes", 34)]
    for key in ("healthcare", "food", "real-estate"):
        s = sectors[key]
        ws = wb.create_sheet(s["name"].replace("&", "and")[:31])
        for c, (h, wdt) in enumerate(cols, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(vertical="center")
            ws.column_dimensions[get_column_letter(c)].width = wdt
        ws.row_dimensions[1].height = 22
        ws.freeze_panes = "D2"
        dv = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        pitch = next(t for t in s["tiers"] if t.get("rec"))
        grp = sorted((r for r in rows if r["launch"] == key), key=lambda r: (-r["score"], -r["reviews"]))
        for i, r in enumerate(grp, 2):
            vals = [i - 1, r["score"], r["name"], r["gmaps_category"], AREA_NAME.get(r["area"], r["area"]),
                    r["neighborhood"], r["rating"], r["reviews"], r["phone"], r["website"], r["website_status"],
                    r["listing_claimed"], opportunity(r), opener(r), f"{pitch['tag']} · {s['headline']}",
                    None, STATUS[0], None, None, None]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=i, column=c, value=v)
                cell.alignment = wrap if c in (3, 13, 14, 15, 20) else Alignment(vertical="top")
                if i <= 26:
                    cell.fill = top_fill
            link = ws.cell(row=i, column=16, value="Open")
            link.hyperlink, link.font = r["maps_url"], Font(color="0079AA", underline="single")
            dv.add(f"Q{i}")
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(grp) + 1}"
        print(f"{s['name']}: {len(grp)} prospects, top: {', '.join(r['name'] for r in grp[:3])}")

    path = OUT / "Islamabad-Prospects.xlsx"
    wb.save(path)
    print(f"saved {path}")


if __name__ == "__main__":
    main()
